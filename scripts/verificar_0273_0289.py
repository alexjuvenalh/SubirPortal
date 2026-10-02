"""Verifica en el portal que las RDs 0273-0289 (17 RDs) estén publicadas.

Mismo patrón que verificar_0268_0272.py: browser con sesión persistente contra
JSON:API autenticado de Drupal (el portal exige sesión para leer el nodo completo).

Fuente de verdad de fecha/numero/resumen: data/resoluciones_pendientes.xlsx
(sheet 'pendientes'). Fids esperados: logs/carga_0273_0289.log.

Checks por RD:
  - titulo exacto (evita falsos positivos del filtro CONTAINS)
  - field_file: >= 1 archivo y contiene TODOS los fids esperados (pdf + plano/anexo)
  - field_tipo_de_norma: target_id == 129
  - field_administrativo: target_id == 17
  - field_fecha_de_aprobacion_de_la_: coincide con fecha_rd del xlsx
  - numero (field_num_resolucion_mat): == numero_resolucion del xlsx
  - body (resumen): prefijo normalizado (60 chars) == resumen_portal del xlsx

Evidencia: logs/verificacion_0273_0289.json (UTF-8).
"""
import asyncio
import datetime
import html
import json
import re
import sys
from pathlib import Path

if sys.platform == "win32":
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "resoluciones_pendientes.xlsx"
OUT_JSON = ROOT / "logs" / "verificacion_0273_0289.json"

from playwright.async_api import async_playwright

NUMEROS = [f"{n:04d}" for n in range(273, 290)]

# Fids reportados por el bot en logs/carga_0273_0289.log (pdf principal + adjuntos)
FIDS_ESPERADOS = {
    "0273": [685045],
    "0274": [685046, 685047],          # + plano
    "0275": [685049],
    "0276": [685052],
    "0277": [685055],
    "0278": [685059],
    "0279": [685062],
    "0280": [685064],
    "0281": [685067],
    "0282": [685069],
    "0283": [685072],
    "0284": [685074],
    "0285": [685077],
    "0286": [685081],
    "0287": [685084, 685086, 685088],  # + anexo + plano
    "0288": [685093],
    "0289": [685096, 685098],          # + plano
}

API = "https://www.ana.gob.pe/jsonapi/node/normatividad"
NUM_KEY = "field_num_resolucion_mat"


def cargar_xlsx():
    """Devuelve {numero_rd: {titulo, fecha(yyyy-mm-dd), numero, resumen, archivo}}."""
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = list(ws.iter_rows(values_only=True))
    cab = [str(c).strip() if c is not None else "" for c in filas[0]]
    idx = {c: i for i, c in enumerate(cab)}
    salida = {}
    for fila in filas[1:]:
        if not fila:
            continue
        get = lambda k: (fila[idx[k]] if k in idx and idx[k] < len(fila) else None)
        titulo = str(get("titulo_sugerido") or "").strip()
        if "N° " not in titulo:
            continue
        numero = titulo.split("N° ")[1].split("-2026")[0]
        fecha = str(get("fecha_rd") or "").strip()  # dd/mm/yyyy
        fecha_iso = ""
        m = re.match(r"^(\d{2})/(\d{2})/(\d{4})$", fecha)
        if m:
            fecha_iso = f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
        salida[numero] = {
            "titulo": titulo,
            "fecha": fecha_iso,
            "numero": str(get("numero_resolucion") or "").strip(),
            "resumen": str(get("resumen_portal") or ""),
            "archivo": str(get("archivo_pdf") or ""),
        }
    return salida


def _target_id(rel_data):
    """Extrae el drupal_internal__target_id de un relationship data (objeto o lista)."""
    if isinstance(rel_data, list):
        return [_target_id(x) for x in rel_data]
    if isinstance(rel_data, dict):
        meta = rel_data.get("meta") or {}
        return meta.get("drupal_internal__target_id")
    return None


def normalizar(texto):
    """Quita HTML, entidades y espacios sobrantes para comparar resúmenes."""
    if not texto:
        return ""
    limpio = re.sub(r"<[^>]+>", " ", str(texto))
    limpio = html.unescape(limpio)
    return re.sub(r"\s+", " ", limpio).strip()


def descubrir_campo_body(attrs):
    """Busca entre los atributos el campo que contiene el resumen (dict con
    'processed' y un value largo). El nombre interno no está en el field map."""
    for key, valor in (attrs or {}).items():
        if not isinstance(valor, dict) or "processed" not in valor:
            continue
        if len(str(valor.get("value", ""))) > 40:
            return key
    return None


def descubrir_campo_numero(attrs, esperado):
    """Si field_num_resolucion_mat no existe, busca otro atributo con el número."""
    if NUM_KEY in (attrs or {}):
        return NUM_KEY
    for key, valor in (attrs or {}).items():
        if isinstance(valor, str) and valor.strip() == str(esperado):
            if "num" in key or "resol" in key:
                return key
    return None


async def main():
    esperados = cargar_xlsx()
    user_data_dir = ROOT / "data" / "browser_profile"
    resultados = []
    campo_body = None
    campo_numero = None

    async with async_playwright() as p:
        context = await p.chromium.launch_persistent_context(
            user_data_dir=str(user_data_dir), headless=False,
            viewport={"width": 1920, "height": 1080}, locale="es-PE")
        page = await context.new_page()

        for numero in NUMEROS:
            esp = esperados.get(numero, {})
            titulo = esp.get("titulo") or f"RD N° {numero}-2026-ANA-AAA.MDD"
            print(f"\n=== {titulo} ===")
            encoded = titulo.replace(" ", "%20")
            # Sin sparse fieldset: Drupal devuelve todos los atributos y relaciones.
            url = (f"{API}?filter[title][operator]=CONTAINS&filter[title][value]={encoded}"
                   f"&page[limit]=5")
            resp = data = None
            ultimo_error = None
            for intento in range(3):
                try:
                    resp = await page.request.get(
                        url, headers={"Accept": "application/vnd.api+json"}, timeout=30000)
                    data = await resp.json()
                    break
                except Exception as e:
                    ultimo_error = str(e).splitlines()[0]
                    print(f"  reintento {intento + 1}/3 tras error: {ultimo_error}")
                    await asyncio.sleep(3)
            if resp is None or data is None:
                print(f"  ERROR consultando API: {ultimo_error}")
                resultados.append({"rd": numero, "status": "ERROR_API",
                                   "detalle": ultimo_error})
                continue

            if resp.status != 200:
                print(f"  HTTP {resp.status}: {str(data)[:200]}")
                resultados.append({"rd": numero, "status": "ERROR_HTTP",
                                   "http": resp.status, "detalle": str(data)[:300]})
                continue

            nodos = data.get("data", [])
            if not nodos:
                print("  NO encontrado en JSON:API")
                resultados.append({"rd": numero, "status": "NO_ENCONTRADO"})
                continue

            exacto = [n for n in nodos if (n.get("attributes") or {}).get("title") == titulo]
            if not exacto:
                print(f"  Solo parcial: {[n['attributes'].get('title') for n in nodos]}")
                resultados.append({"rd": numero, "status": "PARCIAL",
                                   "detalle": [n["attributes"].get("title") for n in nodos]})
                continue

            n = exacto[0]
            attrs = n.get("attributes") or {}
            rel = n.get("relationships") or {}

            if campo_body is None:
                campo_body = descubrir_campo_body(attrs)
                print(f"  Campo de resumen detectado: {campo_body}")
            if campo_numero is None:
                campo_numero = descubrir_campo_numero(attrs, esp.get("numero", ""))
                print(f"  Campo de número detectado: {campo_numero}")

            fecha = attrs.get("field_fecha_de_aprobacion_de_la_")
            numero_real = attrs.get(campo_numero) if campo_numero else None
            ff = rel.get("field_file", {}).get("data") or []
            fids = [x for x in (_target_id(ff) or []) if x is not None]
            ftn = _target_id(rel.get("field_tipo_de_norma", {}).get("data"))
            adm = _target_id(rel.get("field_administrativo", {}).get("data"))
            resumen_esp = normalizar(esp.get("resumen", ""))
            resumen_real = normalizar(attrs.get(campo_body, {}).get("processed", "")) if campo_body else ""

            fecha_esp = esp.get("fecha", "")
            fids_esp = FIDS_ESPERADOS.get(numero, [])

            titulo_ok = True  # ya filtrado por igualdad exacta
            fecha_ok = bool(fecha_esp) and fecha_esp in str(fecha)
            archivos_ok = len(fids) >= 1
            fids_ok = all(f in fids for f in fids_esp)
            numero_ok = (str(numero_real).strip() == str(esp.get("numero", ""))) and bool(esp.get("numero"))
            tipo_ok = ftn == 129
            adm_ok = adm == 17
            resumen_ok = bool(resumen_esp) and resumen_real[:60] == resumen_esp[:60]

            print(f"  fecha:      {fecha} (esperado {fecha_esp}) -> {'OK' if fecha_ok else 'FAIL'}")
            print(f"  archivos:   {fids} (esperados {fids_esp}) -> {'OK' if archivos_ok and fids_ok else 'FAIL'}")
            print(f"  numero:     {numero_real!r} (esperado {esp.get('numero')!r}) -> {'OK' if numero_ok else 'FAIL'}")
            print(f"  tipo_norma: {ftn} (esperado 129) -> {'OK' if tipo_ok else 'FAIL'}")
            print(f"  admin:      {adm} (esperado 17) -> {'OK' if adm_ok else 'FAIL'}")
            print(f"  resumen:    {len(resumen_real)} chars vs {len(resumen_esp)} esperados"
                  f" -> {'OK' if resumen_ok else 'FAIL'}")

            checks = {
                "titulo": titulo_ok,
                "fecha": fecha_ok,
                "archivos": archivos_ok and fids_ok,
                "numero": numero_ok,
                "tipo_de_norma": tipo_ok,
                "administrativo": adm_ok,
                "resumen": resumen_ok,
            }
            fallan = [k for k, v in checks.items() if not v]
            resultados.append({
                "rd": numero,
                "nid": (n.get("id")),
                "status": "OK" if not fallan else "FAIL",
                "campos_fallidos": fallan,
                "fecha": str(fecha),
                "fecha_esperada": fecha_esp,
                "archivos": len(fids),
                "fids": fids,
                "fids_esperados": fids_esp,
                "fids_faltantes": [f for f in fids_esp if f not in fids],
                "numero_real": numero_real,
                "numero_esperado": esp.get("numero"),
                "tipo_de_norma": ftn,
                "administrativo": adm,
                "resumen_chars": len(resumen_real),
                "resumen_esperado_chars": len(resumen_esp),
                "campo_body": campo_body,
                "campo_numero": campo_numero,
            })

        await context.close()

    ok = [r for r in resultados if r["status"] == "OK"]
    print("\n\n======== RESUMEN VERIFICACION 0273-0289 ========")
    print(f"  {'RD':<6} {'ESTADO':<6} {'FECHA':<12} {'FIDS':<26} {'NUM':<6} {'TIPO':<6} {'ADMIN':<6} RESUMEN")
    for r in resultados:
        print(f"  {r['rd']:<6} {r['status']:<6} {str(r.get('fecha', '-'))[:10]:<12} "
              f"{str(r.get('fids', '-')):<26} {str(r.get('numero_real', '-')):<6} "
              f"{str(r.get('tipo_de_norma', '-')):<6} {str(r.get('administrativo', '-')):<6} "
              f"{r.get('resumen_chars', '-')}")
    print(f"\n  Total: {len(resultados)} | OK: {len(ok)} | FAIL: {len(resultados) - len(ok)}")

    OUT_JSON.write_text(
        json.dumps({
            "lote": "0273-0289",
            "generado": datetime.datetime.now().isoformat(timespec="seconds"),
            "fuente_esperado": str(XLSX.relative_to(ROOT)).replace("\\", "/"),
            "fuente_fids": "logs/carga_0273_0289.log",
            "api": API,
            "total": len(resultados),
            "ok": len(ok),
            "fail": len(resultados) - len(ok),
            "resultados": resultados,
        }, ensure_ascii=False, indent=2),
        encoding="utf-8")
    print(f"\nEvidencia guardada: {OUT_JSON}")
    return 0 if len(ok) == len(resultados) else 1


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
