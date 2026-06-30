from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
PDF_DIR = ROOT / "pdfs" / "nuevos"
OUTPUT_XLSX = ROOT / "data" / "resoluciones_pendientes.xlsx"


PREFIX_TO_ALA = {
    "59": "ALA Tahuamanu Madre de Dios",
    "64": "ALA Tambopata - Inambari",
}


SPANISH_MONTHS = {
    "enero": 1,
    "febrero": 2,
    "marzo": 3,
    "abril": 4,
    "mayo": 5,
    "junio": 6,
    "julio": 7,
    "agosto": 8,
    "septiembre": 9,
    "setiembre": 9,
    "octubre": 10,
    "noviembre": 11,
    "diciembre": 12,
}


HEADERS = [
    "archivo_pdf",
    "ruta_pdf",
    "prefijo",
    "ala_detectada",
    "ambito_administrativo",
    "clasificacion_1",
    "clasificacion_2",
    "numero_rd",
    "numero_resolucion",
    "anio_rd",
    "fecha_rd",
    "hora_rd",
    "titulo_sugerido",
    "resumen_portal",
    "articulo_1",
    "articulo_2",
    "articulo_3",
    "articulo_elegido",
]


def normalize_spaces(value: str) -> str:
    return " ".join(value.split())


def extract_text(pdf_path: Path) -> str:
    reader = PdfReader(str(pdf_path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def parse_spanish_date(value: str) -> str:
    match = re.search(r"(\d{1,2})\s+de\s+([a-záéíóúñ]+)\s+de\s+(\d{4})", value, re.I)
    if not match:
        return ""

    day = int(match.group(1))
    month_name = match.group(2).lower()
    year = int(match.group(3))
    month = SPANISH_MONTHS.get(month_name)
    if not month:
        return ""

    return datetime(year, month, day).strftime("%d/%m/%Y")


def extract_resolution_data(pdf_path: Path) -> dict[str, str]:
    text = extract_text(pdf_path)
    filename = pdf_path.name
    prefix = filename.split("-", 1)[0]

    rd_match = re.search(r"([0-9]{4}-[0-9]{4}-ANA-AAA\.MDD)", text, re.I)
    date_match = re.search(r"Tambopata,\s*([^\n]+)", text, re.I)

    numero_rd = rd_match.group(1).upper() if rd_match else ""
    fecha_texto = normalize_spaces(date_match.group(1)) if date_match else ""
    fecha_rd = parse_spanish_date(fecha_texto)

    articulos = extract_articles(text, limit=3)
    anio_rd = numero_rd.split("-")[1] if numero_rd else ""
    numero_resolucion = str(int(numero_rd.split("-", 1)[0])) if numero_rd else ""

    return {
        "archivo_pdf": filename,
        "ruta_pdf": str(pdf_path.resolve()),
        "prefijo": prefix,
        "ala_detectada": PREFIX_TO_ALA.get(prefix, "PENDIENTE_MAPEO"),
        "ambito_administrativo": "AAA Madre de Dios",
        "clasificacion_1": "Resoluciones Emitidas por el ANA",
        "clasificacion_2": "De las Autoridades Administrativas del Agua",
        "numero_rd": numero_rd,
        "numero_resolucion": numero_resolucion,
        "anio_rd": anio_rd,
        "fecha_rd": fecha_rd,
        "hora_rd": "08:00:00 am",
        "titulo_sugerido": f"RD N° {numero_rd}" if numero_rd else "",
        "resumen_portal": articulos[0] if len(articulos) >= 1 else "",
        "articulo_1": articulos[0] if len(articulos) >= 1 else "",
        "articulo_2": articulos[1] if len(articulos) >= 2 else "",
        "articulo_3": articulos[2] if len(articulos) >= 3 else "",
        "articulo_elegido": "1",
    }


def clean_article_ending(article: str) -> str:
    """Asegura que el artículo termine con punto final.

    Si termina en coma, punto y coma, dos puntos, o cualquier otro signo
    de puntuación que no sea . ! ? → lo reemplaza por punto.
    Si no tiene puntuación final → agrega punto.
    """
    article = article.rstrip()
    if not article:
        return article

    last_char = article[-1]
    # Puntuación válida de final de oración: mantener
    if last_char in (".", "!", "?"):
        return article

    # Letra o número: agregar punto
    if last_char.isalnum():
        return article + "."

    # Cualquier otro signo de puntuación (coma, punto y coma, dos puntos, etc.):
    # reemplazar por punto
    return article[:-1] + "."


def remove_table_from_article(article: str) -> str:
    """Corta el texto del artículo justo antes de que empiece un cuadro/tabla.

    Las resoluciones ANA incluyen tablas con coordenadas, vértices, áreas, etc.
    Estos cuadros se ven en el PDF pero no deben subirse al portal institucional.

    Se detectan por frases introductorias como 'cuyo detalle es el siguiente:',
    o por la presencia de marcadores de tabla como 'Vértices', 'Coordenadas UTM',
    'Este (m)', 'Norte (m)', patrones de vértices (v1, v2...), etc.
    """
    # --- Fase 1: frases introductorias de cuadro ---
    # Orden importa: patrones más específicos primero para que no los
    # capture uno más genérico antes.
    table_intro_patterns = [
        r"cuyo\s+detalle.*?(?:siguiente\s+)?cuadro\s*:",
        r"se\s+describen?\s+en\s+el\s+siguiente\s+cuadro\s*:",
        r"cuyo\s+detalle\s+es\s+el\s+siguiente\s*:",
        # "Autorización de ejecución de obras": corta ANTES de "conforme"
        # para que el resumen termine con el nombre del proyecto, sin specs técnicas.
        r"conforme\s+a\s+las?\s+especificaciones\s+t[ée]cnicas\s+del\s+referido\s+proyecto",
        r"conforme\s+al\s+siguiente\s+detalle\s*:",
        r"seg[uú]n\s+el\s+siguiente\s+detalle\s*:",
        r"de\s+acuerdo\s+al\s+siguiente\s+detalle\s*:",
        r"(?:al|el)\s+siguiente\s+detalle\s*:",
        r"detalle\s+siguiente\s*:",
    ]
    for pattern in table_intro_patterns:
        m = re.search(pattern, article, re.IGNORECASE)
        if m:
            return article[: m.start()].rstrip()

    # --- Fase 2: marcadores estructurales de tabla ---
    # Si no hay frase introductoria, buscar patrones que solo aparecen en tablas:
    # - "Vértices" seguido de "Este (m)" o "Norte (m)" en la misma zona
    # - Múltiples referencias a vértices (v1, v2, v3...)
    # - "Coordenadas UTM" seguido de números
    table_markers = [
        r"V[ée]rtices\s+Este\s*\(\s*m\s*\)",       # "Vértices Este (m)"
        r"Coordenadas\s+UTM",                        # "Coordenadas UTM"
        r"UBICACI[ÓO]N\s+GEOGR[ÁA]FICA",             # "UBICACIÓN GEOGRÁFICA"
        r"Cuadro\s+N[°º]?\s*\d+",                   # "Cuadro N° 01", "Cuadro Nro 02"
        r"DATOS\s+DEL\s+ADMINISTRADO",               # tabla de datos en autorizaciones
        r"ESTE\s*\(\s*m\s*\)",                       # columna "Este (m)" en tablas de coordenadas
    ]
    for pattern in table_markers:
        m = re.search(pattern, article, re.IGNORECASE)
        if m:
            return article[: m.start()].rstrip()

    # --- Fase 3: patrón de vértices múltiples ---
    # Si hay 3+ ocurrencias de "v1", "v2", "v3"... con coordenadas, es una tabla
    vertex_matches = re.findall(r"\bv\d+\b", article, re.IGNORECASE)
    if len(vertex_matches) >= 3:
        # Cortar antes del primer vértice
        first_vertex = re.search(r"\bv\d+\b", article, re.IGNORECASE)
        if first_vertex:
            return article[: first_vertex.start()].rstrip()

    return article


# Ordinales usados en resoluciones que usan "SE RESUELVE: PRIMERO. -, SEGUNDO. -..."
_ORDINALES = [
    "PRIMERO", "SEGUNDO", "TERCERO", "CUARTO", "QUINTO",
    "SEXTO", "S[ÉE]PTIMO", "OCTAVO", "NOVENO", "D[ÉE]CIMO",
]


def _extract_articulo_articles(clean_text: str, limit: int) -> list[str]:
    """Formato estándar: ARTÍCULO 1.- ... ARTÍCULO 2.- ...
    
    Soporta variantes:
      - ARTÍCULO 1.-   (guion)
      - ARTÍCULO 1°.-  (símbolo de grado + guion)
      - ARTÍCULO 1º.-  (ordinal masculino + guion)
    """
    # El patrón acepta opcionalmente ° o º entre el número y el .-
    # Solo artículos 1-20: evita capturar referencias legales como "Articulo 257.-"
    art_prefix = r"ART[ÍI]CULO\s+(?:1[0-9]|20|[1-9])\s*[°º]?\s*\.-"
    pattern = re.compile(
        rf"({art_prefix}\s*.*?)(?={art_prefix}|Regístrese|FIRMADO DIGITALMENTE|$)",
        re.I | re.S,
    )
    articles = []
    for match in pattern.finditer(clean_text):
        article = normalize_spaces(match.group(1))
        article = remove_table_from_article(article)
        article = clean_article_ending(article)
        articles.append(article)
        if len(articles) >= limit:
            break
    return articles


def _extract_resuelve_articles(clean_text: str, limit: int) -> list[str]:
    """Formato alternativo: SE RESUELVE: PRIMERO. - ... SEGUNDO. - ..."""
    # Buscar la sección RESUELVE
    resuelve_match = re.search(r"RESUELVE\s*:", clean_text, re.I)
    if not resuelve_match:
        return []

    # Texto desde RESUELVE hasta el final
    post_resuelve = clean_text[resuelve_match.end():]

    # Armar patrón que captura bloques entre ordinales
    # Ej: "PRIMERO. - texto... SEGUNDO. - texto..."
    ordinal_alternation = "|".join(_ORDINALES)
    boundary = rf"(?:{ordinal_alternation})\s*\.\s*-"
    pattern = re.compile(
        rf"({boundary}\s*.*?)(?={boundary}|Regístrese|FIRMADO DIGITALMENTE|$)",
        re.I | re.S,
    )
    articles = []
    for match in pattern.finditer(post_resuelve):
        article = normalize_spaces(match.group(1))
        article = remove_table_from_article(article)
        article = clean_article_ending(article)
        articles.append(article)
        if len(articles) >= limit:
            break
    return articles


def extract_articles(text: str, limit: int = 3) -> list[str]:
    clean_text = remove_authenticity_footers(text)

    # Intentar primero el formato estándar ARTÍCULO X.-
    articles = _extract_articulo_articles(clean_text, limit)
    if articles:
        return articles

    # Si no hay, intentar el formato alternativo RESUELVE: PRIMERO. - / SEGUNDO. -
    return _extract_resuelve_articles(clean_text, limit)


def remove_authenticity_footers(text: str) -> str:
    # pypdf extrae en medio del texto la constancia de copia auténtica que aparece
    # como pie/encabezado de página. La limpiamos para que no ensucie los artículos.
    return re.sub(
        r"Esta es una copia auténtica imprimible.*?ingresando la siguiente clave\s*:\s*[A-Z0-9]+",
        " ",
        text,
        flags=re.I | re.S,
    )


def autosize_columns(ws) -> None:
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            text = str(cell.value or "")
            max_length = max(max_length, min(len(text), 80))
        ws.column_dimensions[column_letter].width = max(max_length + 2, 12)


def write_workbook(rows: list[dict[str, str]]) -> None:
    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "pendientes"
    ws.append(HEADERS)

    for row in rows:
        ws.append([row.get(header, "") for header in HEADERS])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    if rows:
        table_ref = f"A1:{ws.cell(row=len(rows) + 1, column=len(HEADERS)).coordinate}"
        table = Table(displayName="ResolucionesPendientes", ref=table_ref)
        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        table.tableStyleInfo = style
        ws.add_table(table)

    ws.freeze_panes = "A2"
    autosize_columns(ws)
    wb.save(OUTPUT_XLSX)


def es_adjunto(nombre_pdf: str) -> bool:
    """Determina si un archivo es un adjunto (plano o anexo) por su nombre.

    Soporta formatos:
      - RD-XXXX-YYYY-PLANO.pdf          (plano simple)
      - RD-XXXX-YYYY-PLANO-01.pdf       (plano con índice, ej. múltiples planos)
      - RD-XXXX-YYYY-ANEXO.pdf
      - RD-XXXX-YYYY-ANEXO-02.pdf
    """
    return bool(re.search(r"-(PLANO|ANEXO)(-\d+)?\.pdf$", nombre_pdf, re.IGNORECASE))


def main() -> None:
    pdfs = sorted(PDF_DIR.glob("*.pdf"))
    if not pdfs:
        raise SystemExit(f"No se encontraron PDFs en {PDF_DIR}")

    # Filtrar adjuntos — solo procesar RDs principales
    pdfs_principales = [pdf for pdf in pdfs if not es_adjunto(pdf.name)]
    if not pdfs_principales:
        raise SystemExit(f"No se encontraron PDFs de RDs principales en {PDF_DIR} (solo adjuntos)")

    adjuntos = [pdf.name for pdf in pdfs if es_adjunto(pdf.name)]
    if adjuntos:
        print(f"Adjuntos detectados (no se procesan como RDs): {', '.join(adjuntos)}")

    rows = [extract_resolution_data(pdf) for pdf in pdfs_principales]
    # Ordenar por número de resolución (ascendente) para conservar el orden natural
    rows.sort(key=lambda r: int(r["numero_resolucion"]) if r["numero_resolucion"].isdigit() else 9999)
    write_workbook(rows)

    print(f"Planilla generada: {OUTPUT_XLSX}")
    print(f"PDFs procesados: {len(rows)}")


if __name__ == "__main__":
    main()
