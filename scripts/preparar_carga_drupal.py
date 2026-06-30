from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "data" / "resoluciones_pendientes.xlsx"
FIELD_MAP_PATH = ROOT / "config" / "drupal_field_map.json"
OUTPUT_PATH = ROOT / "data" / "carga_drupal_preview.json"
PDFS_DIR = ROOT / "pdfs" / "nuevos"

# Tipos de adjuntos reconocidos (case-insensitive)
# Soporta: -PLANO.pdf, -PLANO1.pdf, -PLANO-01.pdf, -ANEXO.pdf, -ANEXO-02.pdf
ADJUNTO_PATTERNS = [
    re.compile(r"-PLANO(-?\d+)?\.pdf$", re.IGNORECASE),
    re.compile(r"-ANEXO(-?\d+)?\.pdf$", re.IGNORECASE),
]


REQUIRED_COLUMNS = [
    "archivo_pdf",
    "ruta_pdf",
    "ala_detectada",
    "ambito_administrativo",
    "clasificacion_1",
    "clasificacion_2",
    "numero_rd",
    "numero_resolucion",
    "fecha_rd",
    "hora_rd",
    "titulo_sugerido",
    "resumen_portal",
]


def read_rows() -> list[dict[str, str]]:
    wb = load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["pendientes"]
    headers = [cell.value for cell in ws[1]]

    missing = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing:
        raise SystemExit(f"Faltan columnas requeridas en el Excel: {', '.join(missing)}")

    rows: list[dict[str, str]] = []
    for excel_row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        data = {header: normalize_cell(value) for header, value in zip(headers, row)}
        if not data.get("archivo_pdf"):
            continue
        data["_excel_row_number"] = str(excel_row_number)
        rows.append(data)

    return rows


def normalize_cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def to_drupal_date(fecha_rd: str) -> str:
    # Excel generado usa dd/mm/YYYY. Drupal input[type=date] espera YYYY-MM-DD.
    return datetime.strptime(fecha_rd, "%d/%m/%Y").strftime("%Y-%m-%d")


def to_drupal_time(hora_rd: str) -> str:
    # Drupal input[type=time] espera HH:mm:ss en 24 horas.
    normalized = hora_rd.lower().replace(".", "").strip()
    for fmt in ("%I:%M:%S %p", "%I:%M %p", "%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(normalized, fmt).strftime("%H:%M:%S")
        except ValueError:
            pass
    raise ValueError(f"Hora no reconocida: {hora_rd}")


def validate_pdf(path: str) -> str:
    pdf_path = Path(path)
    if not pdf_path.exists():
        raise FileNotFoundError(f"No existe el PDF: {path}")
    if pdf_path.suffix.lower() != ".pdf":
        raise ValueError(f"La ruta no apunta a un PDF: {path}")
    return str(pdf_path.resolve())


def extraer_codigo_base(nombre_pdf: str) -> str:
    """Extrae el código base de una RD desde el nombre del archivo.

    Ejemplos:
        '59-RD-0151-2026-04.pdf'         → '59-RD-0151-2026'
        '59-RD-0151-2026-01-PLANO1.pdf'  → '59-RD-0151-2026'
        '59-RD-0152-2026-04.pdf'         → '59-RD-0152-2026'

    El código base es el prefijo antes del último guion + número(s) + '.pdf'.
    Si el nombre contiene '-PLANO' o '-ANEXO', se remueve esa parte también.
    """
    nombre = Path(nombre_pdf).stem  # sin .pdf
    # Remover sufijo de tipo si existe: -PLANO, -PLANO1, -PLANO-01, -ANEXO, -ANEXO-02, etc.
    nombre = re.sub(r'-(PLANO|ANEXO)(-?\d+)?$', '', nombre, flags=re.IGNORECASE)
    # Remover el último guion + número (la página): -04, -01, -02
    nombre = re.sub(r'-\d+$', '', nombre)
    return nombre


def es_adjunto(nombre_pdf: str) -> bool:
    """Determina si un archivo es un adjunto (plano o anexo) por su nombre."""
    return any(p.search(nombre_pdf) for p in ADJUNTO_PATTERNS)


def detectar_adjuntos(rows: list[dict[str, str]]) -> dict[str, list[str]]:
    """Escanea pdfs/nuevos/ y agrupa adjuntos por código base de RD.

    Retorna un dict: {codigo_base: [adjunto1.pdf, adjunto2.pdf, ...]}
    Solo se consideran adjuntos si existe una RD principal para ese código base.
    Los adjuntos se ordenan por número de página y luego por tipo.
    """
    if not PDFS_DIR.exists():
        return {}

    # Obtener los códigos base de las RDs principales (del Excel)
    codigos_principales: set[str] = set()
    for row in rows:
        principal = row.get("archivo_pdf", "")
        if principal:
            codigos_principales.add(extraer_codigo_base(principal))

    # Escanear pdfs/nuevos/ buscando adjuntos
    adjuntos_por_codigo: dict[str, list[str]] = {codigo: [] for codigo in codigos_principales}

    for pdf_file in sorted(PDFS_DIR.glob("*.pdf")):
        nombre = pdf_file.name
        if not es_adjunto(nombre):
            continue

        codigo = extraer_codigo_base(nombre)
        if codigo in codigos_principales:
            adjuntos_por_codigo[codigo].append(nombre)

    # Ordenar adjuntos: por página, luego alfabéticamente
    for codigo in adjuntos_por_codigo:
        adjuntos_por_codigo[codigo].sort(
            key=lambda n: (
                _extraer_pagina(n),  # ordenar por número de página
                n.lower(),           # desempate alfabético
            )
        )

    return adjuntos_por_codigo


def _extraer_pagina(nombre_pdf: str) -> int:
    """Extrae el número de página de un adjunto para ordenamiento.

    Ej: '59-RD-0151-2026-01-PLANO1.pdf' → 1
        '59-RD-0151-2026-02-ANEXO2.pdf' → 2
    """
    # Buscar el número justo antes de -PLANO o -ANEXO
    m = re.search(r'-(\d+)-(PLANO|ANEXO)', nombre_pdf, re.IGNORECASE)
    if m:
        return int(m.group(1))
    # Fallback: último número antes de .pdf
    parts = Path(nombre_pdf).stem.split('-')
    for part in reversed(parts):
        if part.isdigit():
            return int(part)
    return 99


def build_payload(row: dict[str, str], field_map: dict, adjuntos_map: dict[str, list[str]]) -> dict:
    codigo_base = extraer_codigo_base(row["archivo_pdf"])
    adjuntos = adjuntos_map.get(codigo_base, [])
    # Armar rutas completas para los adjuntos
    adjuntos_rutas = [str((PDFS_DIR / a).resolve()) for a in adjuntos]

    # Determinar qué artículo usar como resumen del portal
    elegido = row.get("articulo_elegido", "1").strip()
    if elegido == "2":
        resumen = row.get("articulo_2", row["resumen_portal"])
    elif elegido == "3":
        resumen = row.get("articulo_3", row["resumen_portal"])
    else:
        resumen = row["resumen_portal"]

    return {
        "excel_row_number": row["_excel_row_number"],
        "archivo_pdf": row["archivo_pdf"],
        "codigo_base": codigo_base,
        "adjuntos": adjuntos,
        "adjuntos_rutas": adjuntos_rutas,
        "url": field_map["url_crear_normatividad"],
        "fields": {
            "titulo": row["titulo_sugerido"],
            "fecha_rd": to_drupal_date(row["fecha_rd"]),
            "hora_rd": to_drupal_time(row["hora_rd"]),
            "numero_resolucion": row["numero_resolucion"],
            "clasificacion_nivel_1": "115",
            "clasificacion_nivel_2": "129",
            "clasificacion_hidden": "129",
            "resumen": resumen,
            "ruta_pdf": validate_pdf(row["ruta_pdf"]),
            "ambito_administrativo_nivel_1": "17",
            "ambito_administrativo_nivel_2": "_none"
        },
    }


def main() -> None:
    if not EXCEL_PATH.exists():
        raise SystemExit(f"No existe el Excel: {EXCEL_PATH}")
    if not FIELD_MAP_PATH.exists():
        raise SystemExit(f"No existe el mapeo Drupal: {FIELD_MAP_PATH}")

    field_map = json.loads(FIELD_MAP_PATH.read_text(encoding="utf-8"))
    rows = read_rows()

    # Detectar adjuntos (planos y anexos) en pdfs/nuevos/
    adjuntos_map = detectar_adjuntos(rows)
    total_adjuntos = sum(len(v) for v in adjuntos_map.values())

    payloads = [build_payload(row, field_map, adjuntos_map) for row in rows]

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(json.dumps(payloads, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Preview de carga generado: {OUTPUT_PATH}")
    print(f"RDs principales: {len(payloads)}")
    if total_adjuntos > 0:
        print(f"Adjuntos detectados: {total_adjuntos}")
        for codigo, adj in adjuntos_map.items():
            if adj:
                print(f"  {codigo}: {', '.join(adj)}")


if __name__ == "__main__":
    main()
